from pathlib import Path
from docx import Document
from docx.shared import Pt,Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_CENTER,TA_JUSTIFY
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer
OUT=Path(r'D:\.pogget\user_storage\u_461180\b9b8b\docs\PatentSubstantiveExam\主动修改准备包_2026-10-17')
paras=[
'申请人在法定主动修改期限内提交本次主动修改文件。申请人此前提交的意见陈述书未附权利要求书全文替换页及修改对照页，未形成完整的权利要求替换文本。本次提交随附《权利要求书全文替换页》和《权利要求书修改对照页》，申请人请求以本次提交的九项权利要求作为后续审查基础。',
'本次修改主要包括：对权利要求1中幕结构及视频合成相关限定进行整理；明确创意版本评价和反馈迭代的适用条件；增加条件解锁DAG状态机引擎的从属限定；整理Shot分镜数据结构、逐幕生成、拖拽排序、后端保存和双视图同源显示关系；进一步明确视频片段时序读取、总时长计算、AI音乐生成降级、静音音轨回退、音频裁剪、混音、视频拼接以及片段生成失败时的兜底机制；对系统权利要求作相应调整。',
'上述修改均能够在原说明书、原权利要求书及附图中找到依据，未引入新的技术内容，未超出原说明书和权利要求书记载的范围。申请人请求国家知识产权局以本次权利要求书全文替换页为基础继续审查。']
text='陈述事项：主动提出修改（根据《中华人民共和国专利法实施细则》第五十七条第一款）\n\n'+'\n\n'.join(paras)+'\n\n附件：1. 权利要求书全文替换页，共2页；2. 权利要求书修改对照页。\n'
(OUT/'03_主动修改意见陈述正文_提交用.txt').write_text(text,encoding='utf-8')
# docx
D=Document(); s=D.sections[0]; s.top_margin=Cm(2.2);s.bottom_margin=Cm(2);s.left_margin=Cm(2.5);s.right_margin=Cm(2.3)
normal=D.styles['Normal'];normal.font.name='宋体';normal.font.size=Pt(11);normal._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体');normal.paragraph_format.line_spacing=1.5;normal.paragraph_format.space_after=Pt(8)
p=D.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run('主动修改意见陈述正文');r.bold=True;r.font.name='黑体';r.font.size=Pt(18);r._element.rPr.rFonts.set(qn('w:eastAsia'),'黑体')
p=D.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run('申请号 2026109280496').bold=True
D.add_paragraph('发明名称：一种基于AI工作流的影视前期制作方法及系统')
D.add_paragraph('陈述事项：主动提出修改（根据《中华人民共和国专利法实施细则》第五十七条第一款）')
p=D.add_paragraph();r=p.add_run('陈述正文');r.bold=True;r.font.name='黑体';r._element.rPr.rFonts.set(qn('w:eastAsia'),'黑体')
for x in paras:D.add_paragraph(x)
p=D.add_paragraph();r=p.add_run('附件清单');r.bold=True;r.font.name='黑体';r._element.rPr.rFonts.set(qn('w:eastAsia'),'黑体')
D.add_paragraph('1. 权利要求书全文替换页，共2页；')
D.add_paragraph('2. 权利要求书修改对照页。')
for sec in D.sections:
 p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run('申请号 2026109280496').font.size=Pt(9)
D.save(OUT/'03_主动修改意见陈述正文_提交用.docx')
# pdf
pdfmetrics.registerFont(TTFont('SimSun',r'C:\Windows\Fonts\simsun.ttc',subfontIndex=0));pdfmetrics.registerFont(TTFont('SimHei',r'C:\Windows\Fonts\simhei.ttf'))
body=ParagraphStyle('body',fontName='SimSun',fontSize=10.5,leading=17,alignment=TA_JUSTIFY,wordWrap='CJK',spaceAfter=6)
title=ParagraphStyle('title',fontName='SimHei',fontSize=16,leading=22,alignment=TA_CENTER,spaceAfter=8)
sub=ParagraphStyle('sub',fontName='SimHei',fontSize=11,leading=18,alignment=TA_CENTER,spaceAfter=10)
h=ParagraphStyle('h',fontName='SimHei',fontSize=12,leading=18,spaceBefore=5,spaceAfter=5)
def footer(c,d):
 c.saveState();c.setFont('SimSun',8);c.drawCentredString(A4[0]/2,12*mm,'申请号 2026109280496');c.restoreState()
story=[Paragraph('主动修改意见陈述正文',title),Paragraph('申请号 2026109280496',sub),Paragraph('发明名称：一种基于AI工作流的影视前期制作方法及系统',body),Paragraph('陈述事项：主动提出修改（根据《中华人民共和国专利法实施细则》第五十七条第一款）',body),Paragraph('陈述正文',h)]
for x in paras:story.append(Paragraph(x,body))
story += [Paragraph('附件清单',h),Paragraph('1. 权利要求书全文替换页，共2页；',body),Paragraph('2. 权利要求书修改对照页。',body)]
doc=SimpleDocTemplate(str(OUT/'03_主动修改意见陈述正文_提交用.pdf'),pagesize=A4,rightMargin=20*mm,leftMargin=22*mm,topMargin=20*mm,bottomMargin=20*mm)
doc.build(story,onFirstPage=footer,onLaterPages=footer)
print('done')
