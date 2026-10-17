from pathlib import Path
import re, shutil
from pypdf import PdfReader
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(r'D:\.pogget\user_storage\u_461180\b9b8b')
OUT=ROOT/'docs'/'PatentSubstantiveExam'/'主动修改准备包_2026-10-17'
OUT.mkdir(parents=True, exist_ok=True)
SRC=Path(r'C:\Users\康泽铭\Downloads')


def extract_pdf(name):
    reader=PdfReader(str(SRC/name))
    return '\n'.join((p.extract_text() or '') for p in reader.pages)

def clean_claims(text):
    lines=[]
    for line in text.splitlines():
        s=line.strip()
        if not s or s in {'权 利 要 求 书','100001','2023.03'} or re.fullmatch(r'[12]',s):
            continue
        lines.append(s)
    s=''.join(lines)
    parts=re.split(r'(?=(?:^|(?<=。))[1-9]\.一种|(?<=。)[2-9]\.根据)', s)
    claims=[]
    for p in parts:
        p=p.strip()
        m=re.match(r'([1-9])\.',p)
        if m: claims.append((int(m.group(1)),p))
    if len(claims)<8:
        claims=[]
        matches=list(re.finditer(r'(?<!\d)([1-9])\.(?=一种|根据)',s))
        for i,m in enumerate(matches):
            claims.append((int(m.group(1)),s[m.start():matches[i+1].start() if i+1<len(matches) else len(s)].strip()))
    return claims

original=clean_claims(extract_pdf('document (1).pdf'))
amended=clean_claims(extract_pdf('MjA2Mjc2NzYw.pdf'))
assert len(original)==8, len(original)
assert len(amended)==9, len(amended)


def set_cell_shading(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)

def set_repeat_header(row):
    trPr=row._tr.get_or_add_trPr(); el=OxmlElement('w:tblHeader'); el.set(qn('w:val'),'true'); trPr.append(el)

def set_margins(cell, top=100, start=110, bottom=100, end=110):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr(); tcMar=tcPr.first_child_found_in('w:tcMar')
    if tcMar is None: tcMar=OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for tag,val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=tcMar.find(qn('w:'+tag))
        if node is None: node=OxmlElement('w:'+tag); tcMar.append(node)
        node.set(qn('w:w'),str(val)); node.set(qn('w:type'),'dxa')

def base_doc(title, subtitle=None):
    d=Document(); sec=d.sections[0]
    sec.top_margin=Cm(2.2); sec.bottom_margin=Cm(2.0); sec.left_margin=Cm(2.4); sec.right_margin=Cm(2.2)
    styles=d.styles
    normal=styles['Normal']; normal.font.name='宋体'; normal.font.size=Pt(11)
    normal._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体')
    normal.paragraph_format.line_spacing=1.45; normal.paragraph_format.space_after=Pt(5)
    for sty,size in [('Title',18),('Heading 1',14),('Heading 2',12)]:
        st=styles[sty]; st.font.name='黑体'; st.font.size=Pt(size); st.font.color.rgb=None
        st._element.rPr.rFonts.set(qn('w:eastAsia'),'黑体')
    p=d.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run(title)
    if subtitle:
        p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(subtitle); r.bold=True
    return d

def add_footer(d):
    for sec in d.sections:
        p=sec.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=p.add_run('申请号 2026109280496'); r.font.name='宋体'; r.font.size=Pt(9)

def save(d,name):
    add_footer(d); path=OUT/name; d.save(path); return path

# 1 claims replacement docx
D=base_doc('权利要求书全文替换页','一种基于AI工作流的影视前期制作方法及系统')
p=D.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('申请号 2026109280496').bold=True
for n,claim in amended:
    p=D.add_paragraph(); p.paragraph_format.first_line_indent=Cm(0); p.paragraph_format.space_after=Pt(8)
    p.add_run(claim)
save(D,'01_权利要求书全文替换页.docx')

# 2 comparison doc
summary={
1:'整理S3幕结构限定；增加视频总时长计算；明确备用AI音乐模型与静音音轨。',
2:'把无条件强制上涨改为针对对应短板实现优化时相对上一版本至少提高5分。',
3:'明确提示词限定评估及修改范围，使反馈迭代针对当前创意方向。',
4:'新增条件解锁DAG状态机、九个工作步骤、完成状态字段及解锁条件。',
5:'由原权利要求4调整为从属项；补入尾帧地址并整理逐幕生成、同源双视图和后端保存关系。',
6:'由原权利要求5调整而来，明确幕摘要用于确定后续片段组织结构。',
7:'由原权利要求6调整为从属项，整理片段读取、总时长、音乐降级、裁剪、混音和拼接。',
8:'由原权利要求7调整而来，明确首帧或实拍参考图、缓慢缩放、兜底标记和参与拼接。',
9:'由原权利要求8调整而来，并补入条件解锁状态机、时长确定及背景音乐降级关系。'}
basis={
1:'原说明书[0005]—[0007]；原权利要求1。',2:'原说明书[0005]、[0008]；原权利要求2。',3:'原说明书[0005]；原权利要求3。',4:'原说明书[0003]、[0004]及图3。',5:'原说明书[0006]及图示说明；原权利要求4。',6:'原说明书[0006]；原权利要求5。',7:'原说明书[0007]、[0008]及图4；原权利要求1、6。',8:'原说明书[0007]、[0008]及图4；原权利要求7。',9:'原说明书[0003]—[0007]；原权利要求8。'}
D=base_doc('权利要求书修改对照页','申请号 2026109280496')
p=D.add_paragraph('说明：本对照页用于说明由原8项权利要求调整为9项权利要求的修改关系。修改后的完整文本以《权利要求书全文替换页》为准。')
t=D.add_table(rows=1,cols=4); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
h=t.rows[0].cells
for i,v in enumerate(['修改后权项','对应原权项','修改内容','原始依据']):
    h[i].text=v; set_cell_shading(h[i],'D9EAF7'); h[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_margins(h[i])
    for r in h[i].paragraphs[0].runs:r.bold=True
set_repeat_header(t.rows[0])
map_orig={1:'1',2:'2',3:'3',4:'说明书新增从属项',5:'4',6:'5',7:'6',8:'7',9:'8'}
for n,_ in amended:
    cells=t.add_row().cells
    vals=[str(n),map_orig[n],summary[n],basis[n]]
    for i,v in enumerate(vals):
        cells[i].text=v; cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_margins(cells[i])
        if i<2: cells[i].paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
for row in t.rows:
    row.cells[0].width=Cm(1.6); row.cells[1].width=Cm(2.5); row.cells[2].width=Cm(8.2); row.cells[3].width=Cm(5.0)
D.add_paragraph('不超范围说明',style='Heading 1')
D.add_paragraph('上述修改均以原申请文件的明确记载为依据，未加入原说明书和原权利要求未记载的模型结构、评分公式、事务机制或其他新的技术内容。修改后的权利要求保护主题仍为原申请公开的AI影视前期制作工作流、结构化分镜及视频合成方案。')
save(D,'02_权利要求书修改对照页.docx')

# 3 statement/body and basis review
D=base_doc('主动修改意见陈述正文及原始依据表','申请号 2026109280496')
D.add_paragraph('发明名称：一种基于AI工作流的影视前期制作方法及系统')
D.add_paragraph('陈述事项：主动提出修改（根据《中华人民共和国专利法实施细则》第五十七条第一款）')
D.add_paragraph('陈述正文',style='Heading 1')
paras=[
'申请人在法定主动修改期限内提交本次主动修改文件。申请人此前提交的意见陈述书未附权利要求书全文替换页及修改对照页，未形成完整的权利要求替换文本。本次提交随附《权利要求书全文替换页》和《权利要求书修改对照页》，申请人请求以本次提交的九项权利要求作为后续审查基础。',
'本次修改主要包括：对权利要求1中幕结构及视频合成相关限定进行整理；明确创意版本评价和反馈迭代的适用条件；增加条件解锁DAG状态机引擎的从属限定；整理Shot分镜数据结构、逐幕生成、拖拽排序、后端保存和双视图同源显示关系；进一步明确视频片段时序读取、总时长计算、AI音乐生成降级、静音音轨回退、音频裁剪、混音、视频拼接以及片段生成失败时的兜底机制；对系统权利要求作相应调整。',
'上述修改均能够在原说明书、原权利要求书及附图中找到依据，未引入新的技术内容，未超出原说明书和权利要求书记载的范围。申请人请求国家知识产权局以本次权利要求书全文替换页为基础继续审查。']
for s in paras:D.add_paragraph(s)
D.add_paragraph('附件清单',style='Heading 1')
for s in ['1. 权利要求书全文替换页，共2页；','2. 权利要求书修改对照页；']: D.add_paragraph(s,style='List Number')
D.add_paragraph('提交前核对项',style='Heading 1')
for s in ['确认已收到进入实质审查阶段通知书，或本次与实质审查请求同时提交；','填写系统中的真实提交日期、通知书发文日及发文序号；','确认电子系统显示三份文件均已上传并进入待提交清单；','正式提交前再次核对局端当前权利要求版本。']: D.add_paragraph(s,style='List Bullet')
D.add_paragraph('逐项原始依据',style='Heading 1')
t=D.add_table(rows=1,cols=3); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
for i,v in enumerate(['权项','主要限定','原申请依据']):
    t.rows[0].cells[i].text=v; set_cell_shading(t.rows[0].cells[i],'D9EAF7'); set_margins(t.rows[0].cells[i])
    for r in t.rows[0].cells[i].paragraphs[0].runs:r.bold=True
set_repeat_header(t.rows[0])
for n,_ in amended:
    cells=t.add_row().cells
    for i,v in enumerate([str(n),summary[n],basis[n]]):
        cells[i].text=v; set_margins(cells[i]); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if i==0: cells[i].paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
save(D,'03_主动修改意见陈述正文及原始依据表.docx')

# copy source replacement PDF with readable name
shutil.copy2(SRC/'MjA2Mjc2NzYw.pdf', OUT/'01_权利要求书全文替换页_原PDF.pdf')
print(OUT)
for p in sorted(OUT.iterdir()): print(p.name,p.stat().st_size)
