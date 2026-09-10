from __future__ import annotations

from pathlib import Path
import html

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Paragraph, PageBreak, Image as PdfImage, Spacer
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "copyright-material"
REAL = OUT / "real-screenshots"
WIDE = OUT / "wide-real-screenshots"
CLEAN_SRC = OUT / "clean-source-screenshots"
SAN = OUT / "sanitized-screenshots"
DOCX = OUT / "AI影视全流程工作流系统_软件使用说明书_脱敏提交版_V1.0.docx"
PDF = OUT / "AI影视全流程工作流系统_软件使用说明书_脱敏提交版_V1.0.pdf"

FONT = ImageFont.truetype(r"C:\Windows\Fonts\msyh.ttc", 22)
FONT_SMALL = ImageFont.truetype(r"C:\Windows\Fonts\msyh.ttc", 18)
FONT_BOLD = ImageFont.truetype(r"C:\Windows\Fonts\msyhbd.ttc", 26)


def cover(draw: ImageDraw.ImageDraw, box, text="", font=FONT, fill=(255, 255, 255), text_fill=(41, 37, 36)):
    draw.rounded_rectangle(box, radius=6, fill=fill)
    if text:
        draw.text((box[0] + 8, box[1] + 4), text, font=font, fill=text_fill)


def save(im: Image.Image, name: str):
    SAN.mkdir(parents=True, exist_ok=True)
    im.save(SAN / name, quality=92)


def redact_sidebar(im: Image.Image, wide=False):
    draw = ImageDraw.Draw(im)
    if wide:
        cover(draw, (52, 398, 238, 548), "示例项目列表", FONT_SMALL)
        cover(draw, (36, 1010, 230, 1066), "测试用户\n点数：1000", FONT_SMALL)
    else:
        cover(draw, (52, 405, 238, 545), "示例项目列表", FONT_SMALL)
        cover(draw, (38, 800, 230, 872), "测试用户\n点数：1000", FONT_SMALL)


def redact_project_title(im: Image.Image, wide=False):
    draw = ImageDraw.Draw(im)
    if wide:
        cover(draw, (460, 155, 1000, 205), "示例项目：博物馆宣传片试行", FONT_BOLD)
    else:
        cover(draw, (315, 150, 760, 190), "示例项目：博物馆宣传片试行", FONT)


def build_sanitized_images():
    # Public pages.
    for src, name in [("00_home.png", "00_home.png"), ("00_login.png", "00_login.png")]:
        Image.open(REAL / src).convert("RGB").save(SAN / name, quality=92)

    # Dashboard, source is the wider stable screenshot.
    im = Image.open(CLEAN_SRC / "02_dashboard.png").convert("RGB")
    draw = ImageDraw.Draw(im)
    cover(draw, (590, 140, 820, 178), "测试用户", FONT_BOLD)
    cover(draw, (82, 610, 342, 820), "示例项目列表", FONT_SMALL)
    cover(draw, (450, 515, 970, 548), "示例项目：博物馆宣传片试行", FONT)
    cover(draw, (1105, 515, 1220, 548), "示例二", FONT)
    cover(draw, (450, 770, 720, 805), "示例项目：城市短片", FONT)
    save(im, "01_dashboard.png")

    # Wide workflow overview.
    im = Image.open(WIDE / "01_workflow.png").convert("RGB")
    redact_sidebar(im, wide=True)
    redact_project_title(im, wide=True)
    save(im, "02_workflow_overview.png")

    # Narrow logged pages: use main-area crops so screenshots are complete where it matters.
    crop_map = [
        ("03_project_overview.png", "03_project_overview.png", "项目总览"),
        ("04_assets.png", "04_assets.png", "资产库"),
        ("05_storyboard.png", "05_storyboard_editor.png", "分镜编辑器"),
        ("06_settings_recharge.png", "06_recharge.png", "充值中心"),
        ("09_step_framework.png", "07_framework.png", "框架搭建"),
        ("10_step_style.png", "08_style.png", "风格统一"),
        ("11_step_character.png", "09_character.png", "人物设计"),
        ("12_step_concept.png", "10_concept.png", "概念图"),
        ("13_step_trailer.png", "11_trailer.png", "宣传片"),
        ("14_step_storyboard_panel.png", "12_storyboard_panel.png", "分镜设计"),
        ("15_step_keyframes.png", "13_keyframes.png", "生成尾帧"),
        ("16_step_direct_video.png", "14_direct_video.png", "直生视频"),
    ]
    for src, out, label in crop_map:
        base = Image.open(REAL / src).convert("RGB")
        # Main content crop removes private sidebar while keeping the real functional panel.
        crop = base.crop((250, 80, base.width - 32, min(base.height - 20, 855)))
        draw = ImageDraw.Draw(crop)
        cover(draw, (62, 72, min(600, crop.width - 20), 112), "示例项目：博物馆宣传片试行", FONT)
        save(crop, out)

    # Admin page: keep table structure but redact emails and names.
    im = Image.open(REAL / "07_admin_users.png").convert("RGB")
    crop = im.crop((88, 80, im.width - 38, min(im.height - 20, 845)))
    draw = ImageDraw.Draw(crop)
    # Table identity columns.
    draw.rectangle((0, 168, 250, crop.height), fill=(255, 255, 255))
    for y in range(175, crop.height - 28, 31):
        draw.text((18, y + 3), "user@example.com", font=FONT_SMALL, fill=(60, 60, 60))
    save(crop, "15_admin_users.png")


SECTIONS = [
    ("系统首页", "00_home.png", "浏览器访问系统首页", "首页展示软件名称、系统定位和进入入口。"),
    ("登录界面", "00_login.png", "首页 → 登录账号", "登录界面用于用户身份验证，保障项目和资产数据隔离。"),
    ("项目仪表盘", "01_dashboard.png", "登录成功 → 仪表盘", "仪表盘展示项目列表、项目状态、统计信息和新建项目入口。"),
    ("工作流看板总览", "02_workflow_overview.png", "项目 → 工作流看板", "工作流看板展示九步影视创作流程、步骤状态和当前步骤入口。"),
    ("项目总览", "03_project_overview.png", "仪表盘 → 项目总览", "项目总览展示项目整体进度、资产概览和工作流入口。"),
    ("资产库", "04_assets.png", "项目 → 资产库", "资产库统一管理图片、视频、音频、文本和参考素材。"),
    ("分镜编辑器", "05_storyboard_editor.png", "项目 → 分镜设计页面", "分镜编辑器支持镜头表格、画面查看、拖拽排序和导出。"),
    ("充值中心", "06_recharge.png", "设置 → 充值", "充值中心展示点数余额、充值记录和凭证提交入口。"),
    ("框架搭建", "07_framework.png", "工作流看板 → 框架搭建", "框架搭建生成故事梗概、角色、幕结构、环境和视觉风格。"),
    ("风格统一", "08_style.png", "工作流看板 → 风格统一", "风格统一生成多组风格图并选定后续生成的统一视觉基准。"),
    ("人物设计", "09_character.png", "工作流看板 → 人物设计", "人物设计基于角色设定和风格基准生成角色概念图。"),
    ("概念图生成", "10_concept.png", "工作流看板 → 概念图", "概念图按幕结构生成关键场景图，形成视觉蓝图。"),
    ("宣传片生成", "11_trailer.png", "工作流看板 → 宣传片", "宣传片模块基于概念图序列生成片段并合成视频。"),
    ("分镜设计步骤", "12_storyboard_panel.png", "工作流看板 → 分镜设计", "分镜设计步骤将故事拆解为镜头序列。"),
    ("生成尾帧", "13_keyframes.png", "工作流看板 → 生成尾帧", "生成尾帧步骤读取首帧并为镜头生成尾帧。"),
    ("直生视频", "14_direct_video.png", "工作流看板 → 直生视频", "直生视频基于首帧或首尾帧生成视频片段并拼接。"),
    ("管理后台", "15_admin_users.png", "管理员入口 → 用户统计", "管理后台提供用户统计、充值审核和运营数据管理。"),
]


def set_run(run, size=None, bold=None, color=None):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def heading(doc, text, level):
    p = doc.add_heading("", level=level)
    r = p.add_run(text)
    set_run(r, size={1: 16, 2: 13}.get(level, 12), bold=True, color="2E74B5")


def para(doc, text, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run(r, bold=bold)
    return p


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def build_docx():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.1)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.1)
    sec.right_margin = Cm(2.1)
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.2
    normal.paragraph_format.space_after = Pt(6)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run("AI 影视全流程工作流系统"), 22, True, "111827")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run("软件使用说明书"), 16, True, "2E74B5")
    for line in ["软件简称：AI Film Flow", "版本号：V1.0", "著作权人：申请人（已脱敏）", "文档类型：软件著作权文档鉴别材料", "完成日期：2026 年 7 月"]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(p.add_run(line), 11)

    doc.add_page_break()
    heading(doc, "一、软件概述", 1)
    para(doc, "AI 影视全流程工作流系统是一款面向影视前期策划与 AI 视频创作的 Web 工作流软件。系统以用户输入的元构思为起点，组织创意扩散、框架搭建、风格统一、人物设计、概念图、宣传片、分镜设计、生成尾帧和直生视频等步骤，辅助用户完成从创意到视频资产的连续化生产。")
    para(doc, "为保护用户隐私和项目资料安全，本文档截图中的用户姓名、邮箱、项目名称、点数等信息已做示例化脱敏处理；界面结构、功能控件、操作路径和软件运行形态均保留真实页面内容。")

    heading(doc, "二、运行环境", 1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "项目"
    table.rows[0].cells[1].text = "说明"
    for c in table.rows[0].cells:
        shade(c, "F2F4F7")
    for k, v in [
        ("客户端", "Chrome、Edge、Firefox 等现代浏览器。"),
        ("服务端", "Next.js 14、React、TypeScript、Prisma、PostgreSQL。"),
        ("存储与队列", "对象存储保存生成资产，队列系统处理异步生成任务。"),
        ("AI 能力", "文本生成、图像生成、图生视频、音乐生成和视频合成。"),
    ]:
        row = table.add_row().cells
        row[0].text = k
        row[1].text = v

    heading(doc, "三、界面截图与操作说明", 1)
    for i, (title, img, path, desc) in enumerate(SECTIONS, 1):
        heading(doc, f"3.{i} {title}", 2)
        para(doc, f"操作路径：{path}", True)
        para(doc, f"功能说明：{desc}")
        if i >= 9:
            para(doc, "说明：本图为该步骤主功能区域截图，用于清晰展示该功能对应的实际运行界面。")
        image_path = SAN / img
        if image_path.exists():
            doc.add_picture(str(image_path), width=Cm(15.8))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_run(cap.add_run(f"图 3-{i} {title}界面"), 9, False, "555555")

    heading(doc, "四、主要功能对应关系", 1)
    for item in [
        "创意扩散：将元构思扩展为多个创作方向，并支持用户选择。",
        "框架搭建：生成故事梗概、角色设定、幕结构、环境设定和视觉风格。",
        "风格统一：生成并选定统一视觉基准。",
        "人物设计与概念图：生成角色和场景视觉资产。",
        "分镜、尾帧与直生视频：将文字和图片资产进一步组织为镜头与视频片段。",
        "资产库、充值和管理后台：提供资产沉淀、点数管理和运营管理能力。",
    ]:
        p = doc.add_paragraph(style="List Bullet")
        set_run(p.add_run(item))
    doc.save(DOCX)


def build_pdf():
    pdfmetrics.registerFont(TTFont("MSYH", r"C:\Windows\Fonts\msyh.ttc"))
    pdfmetrics.registerFont(TTFont("MSYHB", r"C:\Windows\Fonts\msyhbd.ttc"))
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle("TitleCN", fontName="MSYHB", fontSize=22, leading=30, alignment=TA_CENTER, spaceAfter=10))
    styles.add(ParagraphStyle("SubCN", fontName="MSYH", fontSize=11, leading=18, alignment=TA_CENTER, textColor=colors.HexColor("#374151")))
    styles.add(ParagraphStyle("H1CN", fontName="MSYHB", fontSize=16, leading=22, textColor=colors.HexColor("#2E74B5"), spaceBefore=12, spaceAfter=8))
    styles.add(ParagraphStyle("H2CN", fontName="MSYHB", fontSize=13, leading=18, textColor=colors.HexColor("#2E74B5"), spaceBefore=8, spaceAfter=5))
    styles.add(ParagraphStyle("BodyCN", fontName="MSYH", fontSize=10, leading=15, alignment=TA_LEFT, spaceAfter=5))
    styles.add(ParagraphStyle("CapCN", fontName="MSYH", fontSize=8.5, leading=11, alignment=TA_CENTER, textColor=colors.HexColor("#555555"), spaceAfter=7))
    story = [
        Paragraph("AI 影视全流程工作流系统", styles["TitleCN"]),
        Paragraph("软件使用说明书<br/>软件简称：AI Film Flow<br/>版本号：V1.0<br/>著作权人：申请人（已脱敏）<br/>文档类型：软件著作权文档鉴别材料<br/>完成日期：2026 年 7 月", styles["SubCN"]),
        PageBreak(),
        Paragraph("一、软件概述", styles["H1CN"]),
        Paragraph("AI 影视全流程工作流系统是一款面向影视前期策划与 AI 视频创作的 Web 工作流软件。系统以用户输入的元构思为起点，组织创意扩散、框架搭建、风格统一、人物设计、概念图、宣传片、分镜设计、生成尾帧和直生视频等步骤。", styles["BodyCN"]),
        Paragraph("为保护用户隐私和项目资料安全，本文档截图中的用户姓名、邮箱、项目名称、点数等信息已做示例化脱敏处理；界面结构、功能控件、操作路径和软件运行形态均保留真实页面内容。", styles["BodyCN"]),
        Paragraph("二、界面截图与操作说明", styles["H1CN"]),
    ]
    for i, (title, img, path, desc) in enumerate(SECTIONS, 1):
        story.append(Paragraph(f"2.{i} {html.escape(title)}", styles["H2CN"]))
        story.append(Paragraph(f"<b>操作路径：</b>{html.escape(path)}", styles["BodyCN"]))
        story.append(Paragraph(f"<b>功能说明：</b>{html.escape(desc)}", styles["BodyCN"]))
        if i >= 9:
            story.append(Paragraph("说明：本图为该步骤主功能区域截图，用于清晰展示该功能对应的实际运行界面。", styles["BodyCN"]))
        image_path = SAN / img
        if image_path.exists():
            with Image.open(image_path) as im:
                w, h = im.size
            max_w = 16.2 * cm
            img_h = max_w * h / w
            if img_h > 10.2 * cm:
                img_h = 10.2 * cm
                max_w = img_h * w / h
            story.append(PdfImage(str(image_path), width=max_w, height=img_h))
            story.append(Paragraph(f"图 2-{i} {html.escape(title)}界面", styles["CapCN"]))
    story.append(Paragraph("三、主要功能对应关系", styles["H1CN"]))
    for txt in ["创意扩散、框架搭建、风格统一、人物设计、概念图、宣传片、分镜设计、生成尾帧和直生视频构成系统主流程。", "资产库、充值中心和管理后台分别承担资产沉淀、点数管理和运营管理功能。"]:
        story.append(Paragraph(txt, styles["BodyCN"]))
    doc = SimpleDocTemplate(str(PDF), pagesize=A4, leftMargin=2.0 * cm, rightMargin=2.0 * cm, topMargin=1.8 * cm, bottomMargin=1.8 * cm)
    doc.build(story)


if __name__ == "__main__":
    SAN.mkdir(parents=True, exist_ok=True)
    build_sanitized_images()
    build_docx()
    build_pdf()
    print(DOCX)
    print(PDF)
