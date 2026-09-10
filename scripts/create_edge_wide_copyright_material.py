from __future__ import annotations

from pathlib import Path
from statistics import median
from typing import Iterable

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Paragraph, PageBreak, Image as PdfImage, Spacer
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "copyright-material"
RAW = OUT / "edge-wide-raw"
OLD = OUT / "real-screenshots"
SAN = OUT / "edge-wide-sanitized"
DOCX = OUT / "AI影视全流程工作流系统_软件使用说明书_宽屏脱敏提交版_V1.0.docx"
PDF = OUT / "AI影视全流程工作流系统_软件使用说明书_宽屏脱敏提交版_V1.0.pdf"

FONT = ImageFont.truetype(r"C:\Windows\Fonts\msyh.ttc", 18)
FONT_SMALL = ImageFont.truetype(r"C:\Windows\Fonts\msyh.ttc", 15)
FONT_BOLD = ImageFont.truetype(r"C:\Windows\Fonts\msyhbd.ttc", 24)
FONT_TITLE = ImageFont.truetype(r"C:\Windows\Fonts\msyhbd.ttc", 31)

PROJECT = "示例项目：博物馆宣传片试行"
USER = "测试用户"
POINTS = "点数：1000"


def bg(im: Image.Image, box: tuple[int, int, int, int]) -> tuple[int, int, int]:
    x1, y1, x2, y2 = box
    xs = [max(0, min(im.width - 1, x)) for x in (x1 + 3, (x1 + x2) // 2, x2 - 3)]
    ys = [max(0, min(im.height - 1, y)) for y in (y1 + 3, (y1 + y2) // 2, y2 - 3)]
    pts = [im.getpixel((x, y)) for x in xs for y in ys]
    return tuple(int(median([p[i] for p in pts])) for i in range(3))


def fill_text(
    im: Image.Image,
    box: tuple[int, int, int, int],
    text: str,
    font: ImageFont.FreeTypeFont = FONT,
    color: tuple[int, int, int] = (68, 64, 60),
    pad_x: int = 0,
    pad_y: int = 0,
    fill: tuple[int, int, int] | None = None,
) -> None:
    draw = ImageDraw.Draw(im)
    draw.rectangle(box, fill=fill or bg(im, box))
    if text:
        draw.text((box[0] + pad_x, box[1] + pad_y), text, font=font, fill=color)


def redact_sidebar(im: Image.Image) -> None:
    # Project list items.
    fill_text(im, (68, 486, 292, 511), "示例项目一", FONT_SMALL, (120, 113, 108))
    fill_text(im, (68, 522, 292, 548), "示例项目二", FONT_SMALL, (120, 113, 108))
    fill_text(im, (68, 558, 292, 585), "示例项目三", FONT_SMALL, (120, 113, 108))
    fill_text(im, (68, 594, 292, 622), "示例项目四", FONT_SMALL, (120, 113, 108))
    fill_text(im, (68, 630, 292, 658), "示例项目五", FONT_SMALL, (120, 113, 108))
    # Current account footer.
    fill_text(im, (68, 1338, 205, 1372), USER, FONT, (87, 83, 78), fill=(255, 255, 255))
    fill_text(im, (68, 1370, 215, 1398), POINTS, FONT_SMALL, (120, 113, 108), fill=(255, 255, 255))


def redact_project_header(im: Image.Image) -> None:
    fill_text(im, (681, 186, 1188, 226), PROJECT, FONT_TITLE, (41, 37, 36))


def redact_dashboard(im: Image.Image) -> None:
    redact_sidebar(im)
    fill_text(im, (337, 108, 610, 143), "欢迎回来，测试用户", FONT_TITLE, (41, 37, 36))
    cards = [
        (360, 412, "示例项目：博物馆宣传片试行", "创意小故事：用于展示系统项目卡片、当前步骤和资产统计。"),
        (1095, 412, "示例项目：城市宣传片", "示例描述：用于展示项目概览、生成进度和最近操作信息。"),
        (1828, 412, "示例项目：文旅短片", "示例描述：用于展示项目卡片在宽屏栅格中的真实布局。"),
        (360, 613, "示例项目：建筑导览", "示例描述：用于展示多项目列表和不同工作流阶段。"),
        (1095, 613, "示例项目：活动宣传", "示例描述：用于展示项目摘要信息。"),
        (1828, 613, "示例项目：非遗展示", "示例描述：用于展示视频创作流程。"),
        (360, 814, "示例项目：短片练习", "示例描述：用于展示项目入口。"),
        (1095, 814, "示例项目：素材测试", "示例描述：用于展示项目入口。"),
        (1828, 814, "示例项目：视觉测试", "示例描述：用于展示项目入口。"),
        (360, 1016, "示例项目：镜头练习", "示例描述：用于展示项目入口。"),
        (1095, 1016, "示例项目：发货演示", "示例描述：用于展示项目入口。"),
    ]
    for x, y, title, desc in cards:
        fill_text(im, (x, y, x + 655, y + 98), "", FONT, fill=(255, 255, 255))
        fill_text(im, (x, y, x + 430, y + 35), title, FONT_BOLD, (41, 37, 36), fill=(255, 255, 255))
        fill_text(im, (x, y + 43, x + 650, y + 91), desc, FONT, (87, 83, 78), fill=(255, 255, 255))


def redact_admin(im: Image.Image) -> None:
    redact_sidebar(im)
    # Redact visible account/e-mail columns while keeping table structure and row spacing.
    fill_text(im, (330, 315, 535, 1180), "", FONT, fill=(255, 255, 255))
    fill_text(im, (598, 366, 790, 1100), "", FONT, fill=(255, 255, 255))
    for y in [386, 462, 536, 611, 685, 760, 834, 909, 983, 1058]:
        fill_text(im, (604, y, 785, y + 24), "user@example.com", FONT_SMALL, (68, 64, 60), fill=(255, 255, 255))
        fill_text(im, (604, y + 24, 785, y + 46), "测试用户", FONT_SMALL, (120, 113, 108), fill=(255, 255, 255))


def redact_generic_project_page(im: Image.Image) -> None:
    redact_sidebar(im)
    redact_project_header(im)


def build_sanitized_images() -> None:
    SAN.mkdir(parents=True, exist_ok=True)
    # Public pages are not user-specific and keep their existing screenshots.
    for src, out in [("00_home.png", "00_home.png"), ("00_login.png", "00_login.png")]:
        Image.open(OLD / src).convert("RGB").save(SAN / out)

    handlers = {
        "01_dashboard.png": redact_dashboard,
        "06_admin_users.png": redact_admin,
    }
    for p in sorted(RAW.glob("*.png")):
        im = Image.open(p).convert("RGB")
        handler = handlers.get(p.name, redact_generic_project_page)
        handler(im)
        im.save(SAN / p.name)


SECTIONS = [
    ("系统首页", "00_home.png", "浏览器访问系统首页", "首页展示软件名称、系统定位和进入入口。"),
    ("登录界面", "00_login.png", "首页 → 登录账号", "登录界面用于用户身份验证，保障项目和资产数据隔离。"),
    ("项目仪表盘", "01_dashboard.png", "登录成功 → 仪表盘", "仪表盘展示统计数据、项目列表、新建项目入口和最近项目状态。"),
    ("项目总览", "02_project_overview.png", "仪表盘 → 项目总览", "项目总览展示项目整体进度、资产概览和工作流入口。"),
    ("资产库", "03_assets.png", "项目 → 资产库", "资产库统一管理图片、视频、音频、文本和参考素材。"),
    ("分镜编辑器", "04_storyboard_editor.png", "项目 → 分镜设计页面", "分镜编辑器支持镜头表格、画面查看、排序和导出。"),
    ("充值中心", "05_recharge.png", "设置 → 充值", "充值中心展示点数余额、充值记录和凭证提交入口。"),
    ("管理后台", "06_admin_users.png", "管理员入口 → 用户统计", "管理后台提供用户统计、充值审核和运营数据管理。"),
    ("工作流看板总览", "07_workflow_overview.png", "项目 → 工作流看板", "工作流看板展示九步影视创作流程、步骤状态和当前步骤入口。"),
    ("创意扩散", "08_ideation.png", "工作流看板 → 创意扩散", "创意扩散根据项目主题生成多个创意方向。"),
    ("框架搭建", "09_framework.png", "工作流看板 → 框架搭建", "框架搭建生成故事梗概、角色、幕结构、环境和视觉风格。"),
    ("风格统一", "10_style.png", "工作流看板 → 风格统一", "风格统一生成多组风格图并选定后续生成的统一视觉基准。"),
    ("人物设计", "11_character.png", "工作流看板 → 人物设计", "人物设计基于角色设定和风格基准生成角色概念图。"),
    ("概念图", "12_concept.png", "工作流看板 → 概念图", "概念图按幕结构生成关键场景图，形成视觉蓝图。"),
    ("宣传片", "13_trailer.png", "工作流看板 → 宣传片", "宣传片模块基于概念图序列生成片段并合成视频。"),
    ("分镜设计", "14_storyboard_step.png", "工作流看板 → 分镜设计", "分镜设计步骤将故事拆解为镜头序列。"),
    ("生成尾帧", "15_keyframes.png", "工作流看板 → 生成尾帧", "生成尾帧步骤读取首帧并为镜头生成尾帧。"),
    ("直生视频", "16_direct_video.png", "工作流看板 → 直生视频", "直生视频基于首帧或首尾帧生成视频片段并拼接。"),
]


def set_run(run, size=None, bold=None, color=None):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_heading("", level=level)
    r = p.add_run(text)
    set_run(r, size={1: 17, 2: 13}.get(level, 12), bold=True, color="2E74B5")


def add_para(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run(r, size=10.5)


def build_docx() -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width = Cm(29.7)
    sec.page_height = Cm(21)
    sec.top_margin = Cm(1.2)
    sec.bottom_margin = Cm(1.0)
    sec.left_margin = Cm(1.3)
    sec.right_margin = Cm(1.3)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("AI 影视全流程工作流系统")
    set_run(r, 24, True, "000000")
    for line in ["软件使用说明书", "软件简称：AI Film Flow", "版本号：V1.0", "著作权人：申请人（已脱敏）", "文档类型：软件著作权文档鉴别材料", "完成日期：2026 年 7 月"]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line)
        set_run(r, 12, False, "475569")
    doc.add_page_break()

    add_heading(doc, "一、软件概述", 1)
    add_para(doc, "AI 影视全流程工作流系统是一款面向影视前期策划与 AI 视频创作的 Web 工作流软件。系统以用户输入的元构思为起点，组织创意扩散、框架搭建、风格统一、人物设计、概念图、宣传片、分镜设计、生成尾帧和直生视频等步骤。")
    add_para(doc, "为保护用户隐私和项目资料安全，本文档截图中的用户名、邮箱、项目名称、点数等信息已按原界面底色和相近字体做示例化脱敏处理；界面结构、功能控件、页面比例和软件运行形态均保留真实页面内容。")
    doc.add_page_break()

    add_heading(doc, "二、界面截图与操作说明", 1)
    for i, (title, img, path, desc) in enumerate(SECTIONS, start=1):
        add_heading(doc, f"2.{i} {title}", 2)
        add_para(doc, f"操作路径：{path}")
        add_para(doc, f"功能说明：{desc}")
        doc.add_picture(str(SAN / img), width=Cm(26.2))
        cap = doc.add_paragraph(f"图 2-{i} {title}界面")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(cap.runs[0], size=9, color="64748B")
        if i != len(SECTIONS):
            doc.add_page_break()

    doc.add_page_break()
    add_heading(doc, "三、主要功能对应关系", 1)
    rows = [
        ("项目与资产管理", "项目仪表盘、项目总览、资产库", "创建项目、查看统计、维护素材资产。"),
        ("AI 影视工作流", "工作流看板及九个步骤页面", "按步骤完成创意、框架、风格、人物、概念图、宣传片、分镜、尾帧和视频生成。"),
        ("分镜与导出", "分镜编辑器、生成尾帧", "维护镜头信息并支持后续图像、视频生成。"),
        ("运营管理", "充值中心、管理后台", "支持点数管理、充值审核和用户统计。"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ["功能模块", "对应界面", "说明"]):
        cell.text = text
    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text

    doc.save(DOCX)


def build_pdf() -> None:
    pdfmetrics.registerFont(TTFont("MSYH", r"C:\Windows\Fonts\msyh.ttc"))
    pdfmetrics.registerFont(TTFont("MSYHB", r"C:\Windows\Fonts\msyhbd.ttc"))
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("H1CN", parent=styles["Heading1"], fontName="MSYHB", fontSize=18, leading=24, textColor=colors.HexColor("#2E74B5"))
    h2 = ParagraphStyle("H2CN", parent=styles["Heading2"], fontName="MSYHB", fontSize=13, leading=18, textColor=colors.HexColor("#2E74B5"))
    body = ParagraphStyle("BodyCN", parent=styles["BodyText"], fontName="MSYH", fontSize=9.5, leading=15, alignment=TA_LEFT)
    cap = ParagraphStyle("CapCN", parent=styles["BodyText"], fontName="MSYH", fontSize=8.5, leading=12, alignment=TA_CENTER, textColor=colors.HexColor("#64748B"))
    title = ParagraphStyle("TitleCN", parent=styles["Title"], fontName="MSYHB", fontSize=24, leading=34, alignment=TA_CENTER)
    sub = ParagraphStyle("SubCN", parent=styles["BodyText"], fontName="MSYH", fontSize=12, leading=20, alignment=TA_CENTER, textColor=colors.HexColor("#475569"))

    doc = SimpleDocTemplate(str(PDF), pagesize=landscape(A4), rightMargin=1.2 * cm, leftMargin=1.2 * cm, topMargin=1.0 * cm, bottomMargin=0.8 * cm)
    story = [
        Spacer(1, 2.5 * cm),
        Paragraph("AI 影视全流程工作流系统", title),
        Spacer(1, 0.35 * cm),
        Paragraph("软件使用说明书<br/>软件简称：AI Film Flow<br/>版本号：V1.0<br/>著作权人：申请人（已脱敏）<br/>文档类型：软件著作权文档鉴别材料<br/>完成日期：2026 年 7 月", sub),
        PageBreak(),
        Paragraph("一、软件概述", h1),
        Paragraph("AI 影视全流程工作流系统是一款面向影视前期策划与 AI 视频创作的 Web 工作流软件。系统以用户输入的元构思为起点，组织创意扩散、框架搭建、风格统一、人物设计、概念图、宣传片、分镜设计、生成尾帧和直生视频等步骤。", body),
        Spacer(1, 0.25 * cm),
        Paragraph("为保护用户隐私和项目资料安全，本文档截图中的用户名、邮箱、项目名称、点数等信息已按原界面底色和相近字体做示例化脱敏处理；界面结构、功能控件、页面比例和软件运行形态均保留真实页面内容。", body),
        PageBreak(),
        Paragraph("二、界面截图与操作说明", h1),
    ]
    for i, (section_title, img, path, desc) in enumerate(SECTIONS, start=1):
        story += [
            Paragraph(f"2.{i} {section_title}", h2),
            Paragraph(f"操作路径：{path}", body),
            Paragraph(f"功能说明：{desc}", body),
            Spacer(1, 0.15 * cm),
            PdfImage(str(SAN / img), width=26.9 * cm, height=14.75 * cm),
            Paragraph(f"图 2-{i} {section_title}界面", cap),
        ]
        if i != len(SECTIONS):
            story.append(PageBreak())
    story += [
        PageBreak(),
        Paragraph("三、主要功能对应关系", h1),
        Paragraph("创意扩散、框架搭建、风格统一、人物设计、概念图、宣传片、分镜设计、生成尾帧和直生视频构成系统主流程。资产库、充值中心和管理后台分别承担资产沉淀、点数管理和运营管理功能。", body),
    ]
    doc.build(story)


if __name__ == "__main__":
    build_sanitized_images()
    build_docx()
    build_pdf()
    print(DOCX)
    print(PDF)
