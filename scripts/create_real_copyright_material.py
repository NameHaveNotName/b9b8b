from __future__ import annotations

from pathlib import Path
import html

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Paragraph, PageBreak, Image as PdfImage
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "copyright-material"
SCREEN_DIR = OUT / "real-screenshots"
DOCX_OUT = OUT / "AI影视全流程工作流系统_补正文档鉴别材料_真实界面版_V1.0.docx"
PDF_OUT = OUT / "AI影视全流程工作流系统_补正文档鉴别材料_真实界面版_V1.0.pdf"


SECTIONS = [
    {
        "title": "系统首页",
        "path": "浏览器访问系统首页",
        "image": "00_home.png",
        "desc": "首页展示软件名称、软件定位和进入系统的操作入口，是用户访问 AI 影视全流程工作流系统的起始界面。",
        "steps": ["用户在浏览器地址栏输入系统访问地址。", "系统展示软件名称、简介和开始使用入口。", "用户点击开始使用或登录账号进入后续操作。"],
    },
    {
        "title": "登录界面",
        "path": "首页 → 登录账号",
        "image": "00_login.png",
        "desc": "登录界面用于用户身份验证，保障项目数据、资产数据和后台权限隔离。",
        "steps": ["用户输入邮箱和密码。", "系统调用认证服务校验用户身份。", "登录成功后进入项目仪表盘或原目标项目页面。"],
    },
    {
        "title": "项目仪表盘",
        "path": "登录成功 → 仪表盘",
        "image": "02_dashboard.png",
        "desc": "项目仪表盘集中展示用户创建的项目、项目进度、更新时间和项目入口，支持进入已有项目或创建新项目。",
        "steps": ["系统读取当前登录用户的项目列表。", "页面以卡片方式展示项目名称、进度和状态。", "用户点击项目卡片进入项目总览或工作流看板。"],
    },
    {
        "title": "项目总览",
        "path": "仪表盘 → 选择项目 → 项目总览",
        "image": "03_project_overview.png",
        "desc": "项目总览页面展示当前项目的基础信息、整体进度、最近资产和工作流入口，帮助用户把握项目全局状态。",
        "steps": ["用户从仪表盘进入指定项目。", "系统展示项目标题、工作流完成进度和资产概览。", "用户可进入工作流、资产库或其他项目功能。"],
    },
    {
        "title": "工作流看板与创意扩散",
        "path": "项目总览 → 工作流看板 → 创意扩散",
        "image": "01_workflow.png",
        "desc": "工作流看板以 9 个步骤组织项目生产流程。创意扩散步骤将原始元构思扩展为多个创意方向，并允许用户选择后续使用的方向。",
        "steps": ["用户进入工作流看板。", "系统展示九步工作流状态、参考素材和当前步骤内容。", "用户在创意扩散步骤查看生成方向、标签和描述，并选择合适方向。"],
    },
    {
        "title": "框架搭建",
        "path": "工作流看板 → 框架搭建",
        "image": "09_step_framework.png",
        "desc": "框架搭建步骤生成故事梗概、角色设定、幕结构、环境设定和视觉风格，是后续图像与视频生成的结构化基础。",
        "steps": ["用户点击框架搭建步骤。", "系统展示已生成的故事框架内容。", "用户可编辑文本字段，并使用角色、故事或环境深化能力补充细节。"],
    },
    {
        "title": "风格统一",
        "path": "工作流看板 → 风格统一",
        "image": "10_step_style.png",
        "desc": "风格统一步骤生成多组视觉风格样图，用户选择其中一组作为后续人物、概念图和视频生成的统一视觉基准。",
        "steps": ["用户进入风格统一步骤。", "系统展示多张风格样图、模型标记和风格说明。", "用户选择符合项目基调的风格图并设置为后续生成基准。"],
    },
    {
        "title": "人物设计",
        "path": "工作流看板 → 人物设计",
        "image": "11_step_character.png",
        "desc": "人物设计步骤基于角色设定和统一风格图生成角色概念图，形成后续场景图和视频生成的角色参考资产。",
        "steps": ["系统读取框架中的角色列表。", "用户查看角色设定、提示词和已生成角色图。", "用户可对单个角色重新生成或调整提示词。"],
    },
    {
        "title": "概念图生成",
        "path": "工作流看板 → 概念图",
        "image": "12_step_concept.png",
        "desc": "概念图步骤按幕结构生成关键场景图，并结合风格参考和角色参考形成项目视觉蓝图。",
        "steps": ["用户进入概念图步骤。", "系统按幕展示场景提示词和生成结果。", "用户可确认、预览或对单张概念图重新生成。"],
    },
    {
        "title": "宣传片生成",
        "path": "工作流看板 → 宣传片",
        "image": "13_step_trailer.png",
        "desc": "宣传片步骤基于概念图序列生成多个视频片段，并支持合成为完整宣传片。",
        "steps": ["系统读取概念图和框架内容。", "用户为片段生成视频提示词并逐段生成视频。", "全部片段完成后系统合成宣传片并保存为视频资产。"],
    },
    {
        "title": "分镜设计面板",
        "path": "工作流看板 → 分镜设计",
        "image": "14_step_storyboard_panel.png",
        "desc": "分镜设计步骤将故事拆解为镜头序列，连接文字脚本和后续尾帧、视频生成流程。",
        "steps": ["用户进入分镜设计步骤。", "系统展示镜头列表、分镜模式和生成状态。", "用户可进入分镜编辑器进一步调整镜头信息。"],
    },
    {
        "title": "分镜编辑器",
        "path": "项目 → 分镜设计页面",
        "image": "05_storyboard.png",
        "desc": "分镜编辑器提供表格和画布等视图，支持镜头行内编辑、拖拽排序、画面查看和分镜导出。",
        "steps": ["用户打开分镜设计页面。", "系统展示镜头编号、画面描述、运镜、时长和参考画面。", "用户可编辑、排序并导出 JSON 或 Excel 分镜文件。"],
    },
    {
        "title": "生成尾帧",
        "path": "工作流看板 → 生成尾帧",
        "image": "15_step_keyframes.png",
        "desc": "生成尾帧步骤读取分镜中的首帧作为只读输入，为镜头生成尾帧，供直生视频使用。",
        "steps": ["系统读取分镜步骤的首帧信息。", "用户查看尾帧提示词和生成状态。", "系统逐个镜头生成尾帧并保存结果。"],
    },
    {
        "title": "直生视频",
        "path": "工作流看板 → 直生视频",
        "image": "16_step_direct_video.png",
        "desc": "直生视频步骤基于首帧或首尾帧调用图生视频模型生成视频片段，并可拼接为完整视频。",
        "steps": ["用户进入直生视频步骤。", "系统展示视频模型、片段状态和生成入口。", "用户批量生成片段并在完成后拼接视频。"],
    },
    {
        "title": "项目资产库",
        "path": "项目 → 资产库",
        "image": "04_assets.png",
        "desc": "资产库统一管理项目生成的图片、视频、音频、文本和参考素材，支持预览和按类型查看。",
        "steps": ["用户进入项目资产库。", "系统读取项目关联资产。", "用户按资产类型查看、预览或下载项目成果。"],
    },
    {
        "title": "个人中心与充值",
        "path": "设置 → 充值",
        "image": "06_settings_recharge.png",
        "desc": "充值页面展示点数余额、充值记录和凭证上传入口，用于支撑 AI 生成任务的点数消耗。",
        "steps": ["用户进入充值页面。", "系统展示当前点数和充值记录。", "用户提交充值金额和凭证，等待管理员审核。"],
    },
    {
        "title": "管理后台",
        "path": "管理员入口 → 用户统计",
        "image": "07_admin_users.png",
        "desc": "管理后台提供用户统计、充值审核、请求统计和项目查看能力，用于系统运营管理。",
        "steps": ["管理员进入后台页面。", "系统展示用户列表、点数和操作入口。", "管理员可查看用户项目、调整点数或审核充值订单。"],
    },
]


def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_heading(doc, text, level):
    p = doc.add_heading("", level=level)
    r = p.add_run(text)
    set_run_font(r, size={1: 16, 2: 13, 3: 12}.get(level, 12), bold=True, color="2E74B5" if level < 3 else "1F4D78")
    return p


def add_para(doc, text, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run_font(r, bold=bold)
    return p


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def build_docx():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.2
    normal.paragraph_format.space_after = Pt(6)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("AI 影视全流程工作流系统"), size=22, bold=True, color="111827")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("软件著作权补正文档鉴别材料（真实运行界面版）"), size=15, bold=True, color="2E74B5")
    for line in ["软件简称：AI Film Flow", "版本号：V1.0", "著作权人：康泽铭", "对应补正流水号：2026R11L2059409", "完成日期：2026 年 7 月"]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(p.add_run(line), size=11)

    doc.add_page_break()
    add_heading(doc, "一、补正说明", 1)
    add_para(doc, "本材料根据软件登记补正通知书要求编制，重点补充 AI 影视全流程工作流系统运行过程中全部主要功能对应的真实界面截图，并配合文字介绍软件相关操作。")
    add_para(doc, "截图来源为软件线上运行环境中的真实页面，覆盖系统首页、登录、项目仪表盘、项目总览、九步工作流、资产库、充值页面和管理后台。")

    add_heading(doc, "二、软件基本情况", 1)
    t = doc.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    t.rows[0].cells[0].text = "项目"
    t.rows[0].cells[1].text = "内容"
    for c in t.rows[0].cells:
        shade(c, "F2F4F7")
    for k, v in [
        ("软件全称", "AI 影视全流程工作流系统"),
        ("软件简称", "AI Film Flow"),
        ("版本号", "V1.0"),
        ("软件类型", "Web 应用软件 / AI 影视创作工作流系统"),
        ("主要用途", "辅助用户从元构思开始，完成创意扩散、框架搭建、风格统一、人物设计、概念图、宣传片、分镜、尾帧和直生视频生成。"),
        ("运行方式", "用户通过浏览器访问系统，无需安装客户端。"),
    ]:
        row = t.add_row().cells
        row[0].text = k
        row[1].text = v

    add_heading(doc, "三、真实界面截图与操作说明", 1)
    for idx, item in enumerate(SECTIONS, 1):
        add_heading(doc, f"3.{idx} {item['title']}", 2)
        add_para(doc, f"操作路径：{item['path']}", bold=True)
        add_para(doc, f"功能说明：{item['desc']}")
        add_para(doc, "操作步骤：")
        for step in item["steps"]:
            p = doc.add_paragraph(style="List Number")
            set_run_font(p.add_run(step))
        image_path = SCREEN_DIR / item["image"]
        if image_path.exists():
            doc.add_picture(str(image_path), width=Cm(15.8))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_run_font(cap.add_run(f"图 3-{idx} {item['title']}真实运行界面"), size=9, color="555555")

    add_heading(doc, "四、功能与技术特点对应说明", 1)
    for text in [
        "全流程工作流：系统通过九步工作流组织影视前期制作流程，步骤状态可视化显示，支持已完成、进行中、失败、跳过等状态。",
        "AI 生成与人工确认：创意、框架、风格图、角色图、概念图、视频片段等内容均支持 AI 生成后由用户预览、选择、确认或重新生成。",
        "统一视觉基准：风格统一步骤生成并选定统一风格图，后续人物设计、概念图和视频生成均围绕该视觉基准展开。",
        "资产沉淀与复用：系统将生成的图片、视频、音频、文本和参考素材保存至项目资产库，便于后续查看、复用和导出。",
        "运营管理能力：系统提供点数充值、充值审核、用户统计、请求统计等后台能力，适用于多人和商业化使用场景。",
    ]:
        p = doc.add_paragraph(style="List Bullet")
        set_run_font(p.add_run(text))

    add_heading(doc, "五、补正结论", 1)
    add_para(doc, "本补正文档已针对通知书指出的文档鉴别材料缺陷，补充软件运行全部主要功能对应的真实界面截图，并配套说明相关操作路径和操作过程。材料内容能够体现软件主要功能、技术特点和运行界面，可作为 AI 影视全流程工作流系统 V1.0 的软件著作权补正文档鉴别材料提交。")
    doc.save(DOCX_OUT)


def build_pdf():
    pdfmetrics.registerFont(TTFont("MSYH", r"C:\Windows\Fonts\msyh.ttc"))
    pdfmetrics.registerFont(TTFont("MSYHB", r"C:\Windows\Fonts\msyhbd.ttc"))
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle("TitleCN", fontName="MSYHB", fontSize=22, leading=30, alignment=TA_CENTER, spaceAfter=10))
    styles.add(ParagraphStyle("SubCN", fontName="MSYH", fontSize=11, leading=18, alignment=TA_CENTER, textColor=colors.HexColor("#374151")))
    styles.add(ParagraphStyle("H1CN", fontName="MSYHB", fontSize=16, leading=22, textColor=colors.HexColor("#2E74B5"), spaceBefore=12, spaceAfter=8))
    styles.add(ParagraphStyle("H2CN", fontName="MSYHB", fontSize=13, leading=19, textColor=colors.HexColor("#2E74B5"), spaceBefore=10, spaceAfter=6))
    styles.add(ParagraphStyle("BodyCN", fontName="MSYH", fontSize=10.2, leading=15.5, alignment=TA_LEFT, spaceAfter=5))
    styles.add(ParagraphStyle("CapCN", fontName="MSYH", fontSize=8.6, leading=11, alignment=TA_CENTER, textColor=colors.HexColor("#555555"), spaceAfter=7))
    story = [
        Paragraph("AI 影视全流程工作流系统", styles["TitleCN"]),
        Paragraph("软件著作权补正文档鉴别材料（真实运行界面版）<br/>软件简称：AI Film Flow<br/>版本号：V1.0<br/>著作权人：康泽铭<br/>对应补正流水号：2026R11L2059409<br/>完成日期：2026 年 7 月", styles["SubCN"]),
        PageBreak(),
        Paragraph("一、补正说明", styles["H1CN"]),
        Paragraph("本材料根据软件登记补正通知书要求编制，重点补充 AI 影视全流程工作流系统运行过程中全部主要功能对应的真实界面截图，并配合文字介绍软件相关操作。", styles["BodyCN"]),
        Paragraph("二、真实界面截图与操作说明", styles["H1CN"]),
    ]
    for idx, item in enumerate(SECTIONS, 1):
        story.append(Paragraph(f"2.{idx} {html.escape(item['title'])}", styles["H2CN"]))
        story.append(Paragraph(f"<b>操作路径：</b>{html.escape(item['path'])}", styles["BodyCN"]))
        story.append(Paragraph(f"<b>功能说明：</b>{html.escape(item['desc'])}", styles["BodyCN"]))
        story.append(Paragraph("<b>操作步骤：</b>" + "；".join(html.escape(s) for s in item["steps"]) + "。", styles["BodyCN"]))
        image_path = SCREEN_DIR / item["image"]
        if image_path.exists():
            story.append(PdfImage(str(image_path), width=16.2 * cm, height=10.1 * cm))
            story.append(Paragraph(f"图 2-{idx} {html.escape(item['title'])}真实运行界面", styles["CapCN"]))
    story.append(Paragraph("三、补正结论", styles["H1CN"]))
    story.append(Paragraph("本补正文档已针对通知书指出的文档鉴别材料缺陷，补充软件运行全部主要功能对应的真实界面截图，并配套说明相关操作路径和操作过程。材料内容能够体现软件主要功能、技术特点和运行界面，可作为 AI 影视全流程工作流系统 V1.0 的软件著作权补正文档鉴别材料提交。", styles["BodyCN"]))
    doc = SimpleDocTemplate(str(PDF_OUT), pagesize=A4, leftMargin=2.0 * cm, rightMargin=2.0 * cm, topMargin=1.8 * cm, bottomMargin=1.8 * cm)
    doc.build(story)


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    build_docx()
    build_pdf()
    print(DOCX_OUT)
    print(PDF_OUT)
