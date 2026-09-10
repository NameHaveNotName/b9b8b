from __future__ import annotations

import json
import re
import subprocess
import sys
from pathlib import Path

import pdfplumber
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
PDF_SOURCE = Path(r"C:\Users\康泽铭\Downloads\dfebbb3d-c7f1-41af-8640-5b309d054800.pdf")
OUT_DIR = ROOT / "output" / "copyright-material"
SCREEN_DIR = OUT_DIR / "real-screenshots"
OUT_DOCX = OUT_DIR / "AI影视全流程工作流系统_软件使用说明书_真实界面配图版_V1.0.docx"

SOFTWARE_NAME = "AI 影视全流程工作流系统"
VERSION = "V1.0"
COPYRIGHT_OWNER = "康泽铭"
COMPLETION_DATE = "2026 年 6 月 7 日"


def text_from_pdf() -> str:
    with pdfplumber.open(PDF_SOURCE) as pdf:
        return "\n".join(page.extract_text() or "" for page in pdf.pages)


PDF_TEXT = text_from_pdf()


def pick_pdf_paragraph(pattern: str, fallback: str, max_len: int = 900) -> str:
    match = re.search(pattern, PDF_TEXT, re.S)
    if not match:
        return fallback
    text = re.sub(r"\s+", " ", match.group(1)).strip()
    return text[:max_len]


INTRO_TEXT = pick_pdf_paragraph(
    r"1\.1\s*软件简介\s*(.*?)(?:1\.2\s*运行环境)",
    f"{SOFTWARE_NAME}是一款面向影视工业化生产的 AI 驱动在线工作流系统。系统以元构思为起点，通过多步骤人工智能协作，帮助用户完成从创意扩散、剧本框架、视觉风格、人物设计、概念图、分镜设计到宣传片与直出视频生成的完整影视前期制作流程。",
)

ENV_TEXT = pick_pdf_paragraph(
    r"1\.2\s*运行环境\s*(.*?)(?:1\.3\s*主要功能模块)",
    "客户端支持 Chrome、Firefox、Edge 等现代 Web 浏览器。服务端基于 Next.js、React、TypeScript、Prisma、PostgreSQL、对象存储、异步任务队列与 ffmpeg 等技术构建，支持在线化访问和部署。",
)

FEATURE_TEXT = pick_pdf_paragraph(
    r"1\.4\s*软件特点\s*(.*?)(?:1\.5\s*使用场景)",
    "系统覆盖影视前期创作全流程，支持 AI 自动生成与人工确认结合，具备视觉风格一致性、分段式视频管理、多模型调度、点数与充值管理、Web 化部署和管理后台等特点。",
)


SECTIONS = [
    {
        "title": "系统首页",
        "path": "浏览器访问系统首页",
        "image": "00_home.png",
        "desc": "系统首页展示软件名称、产品定位、核心能力简介和进入系统的操作入口，是用户访问 AI 影视全流程工作流系统的起始界面。",
        "steps": ["用户在浏览器地址栏输入系统访问地址。", "系统展示软件名称、简介、核心能力和开始使用入口。", "用户点击开始使用或登录账号进入后续操作。"],
    },
    {
        "title": "登录界面",
        "path": "首页 - 登录账号",
        "image": "00_login.png",
        "desc": "登录界面用于用户身份验证，保证项目数据、生成资产和后台权限隔离。",
        "steps": ["用户输入邮箱和密码。", "系统调用认证服务校验用户身份。", "登录成功后进入项目仪表盘或目标项目页面。"],
    },
    {
        "title": "项目仪表盘",
        "path": "登录成功 - 仪表盘",
        "image": "02_dashboard.png",
        "desc": "项目仪表盘集中展示用户创建的项目、项目进度、更新时间和项目入口，支持进入已有项目或创建新项目。",
        "steps": ["系统读取当前登录用户的项目列表。", "页面以卡片方式展示项目名称、进度和状态。", "用户点击项目卡片进入项目总览或工作流看板。"],
    },
    {
        "title": "项目总览",
        "path": "仪表盘 - 选择项目 - 项目总览",
        "image": "03_project_overview.png",
        "desc": "项目总览页面展示当前项目的基础信息、整体进度、最近资产和工作流入口，帮助用户把握项目全局状态。",
        "steps": ["用户从仪表盘进入指定项目。", "系统展示项目标题、工作流完成进度和资产概览。", "用户可进入工作流、资产库或其他项目功能。"],
    },
    {
        "title": "工作流看板与创意扩散",
        "path": "项目总览 - 工作流看板 - 创意扩散",
        "image": "01_workflow.png",
        "desc": "工作流看板以 9 个步骤组织项目生产流程。创意扩散步骤将原始元构思扩展为多个创意方向，并允许用户选择后续使用的方向。",
        "steps": ["用户进入工作流看板。", "系统展示九步工作流状态、参考素材和当前步骤内容。", "用户在创意扩散步骤查看生成方向、标签和描述，并选择合适方向。"],
    },
    {
        "title": "框架搭建",
        "path": "工作流看板 - 框架搭建",
        "image": "09_step_framework.png",
        "desc": "框架搭建步骤生成故事梗概、角色设定、幕结构、环境设定和视觉风格，是后续图像与视频生成的结构化基础。",
        "steps": ["用户点击框架搭建步骤。", "系统展示已生成的故事框架内容。", "用户可编辑文本字段，并使用角色、故事或环境深化能力补充细节。"],
    },
    {
        "title": "风格统一",
        "path": "工作流看板 - 风格统一",
        "image": "10_step_style.png",
        "desc": "风格统一步骤生成多组视觉风格样图，用户选择其中一组作为后续人物、概念图和视频生成的统一视觉基准。",
        "steps": ["用户进入风格统一步骤。", "系统展示多张风格样图、模型标记和风格说明。", "用户选择符合项目基调的风格图并设为后续生成基准。"],
    },
    {
        "title": "人物设计",
        "path": "工作流看板 - 人物设计",
        "image": "11_step_character.png",
        "desc": "人物设计步骤基于角色设定和统一风格图生成角色概念图，形成后续场景图和视频生成的角色参考资产。",
        "steps": ["系统读取框架中的角色列表。", "用户查看角色设定、提示词和已生成角色图。", "用户可对单个角色重新生成或调整提示词。"],
    },
    {
        "title": "概念图生成",
        "path": "工作流看板 - 概念图",
        "image": "12_step_concept.png",
        "desc": "概念图步骤按幕结构生成关键场景图，并结合风格参考和角色参考形成项目视觉蓝图。",
        "steps": ["用户进入概念图步骤。", "系统按幕展示场景提示词和生成结果。", "用户可确认、预览或对单张概念图重新生成。"],
    },
    {
        "title": "宣传片生成",
        "path": "工作流看板 - 宣传片",
        "image": "13_step_trailer.png",
        "desc": "宣传片步骤基于概念图序列生成多个视频片段，并支持合成为完整宣传片。",
        "steps": ["系统读取概念图和框架内容。", "用户为片段生成视频提示词并逐段生成视频。", "全部片段完成后系统合成宣传片并保存为视频资产。"],
    },
    {
        "title": "分镜设计面板",
        "path": "工作流看板 - 分镜设计",
        "image": "14_step_storyboard_panel.png",
        "desc": "分镜设计步骤将故事拆解为镜头序列，连接文字脚本和后续尾帧、视频生成流程。",
        "steps": ["用户进入分镜设计步骤。", "系统展示镜头列表、分镜模式和生成状态。", "用户可进入分镜编辑器进一步调整镜头信息。"],
    },
    {
        "title": "分镜编辑器",
        "path": "项目 - 分镜设计页面",
        "image": "05_storyboard.png",
        "desc": "分镜编辑器提供表格和画布等视图，支持镜头行内编辑、拖拽排序、画面查看和分镜导出。",
        "steps": ["用户打开分镜设计页面。", "系统展示镜头编号、画面描述、运镜、时长和参考画面。", "用户可编辑、排序并导出 JSON 或 Excel 分镜文件。"],
    },
    {
        "title": "生成尾帧",
        "path": "工作流看板 - 生成尾帧",
        "image": "15_step_keyframes.png",
        "desc": "生成尾帧步骤读取分镜中的首帧作为输入，为镜头生成尾帧，供直生视频使用。",
        "steps": ["系统读取分镜步骤的首帧信息。", "用户查看尾帧提示词和生成状态。", "系统逐个镜头生成尾帧并保存结果。"],
    },
    {
        "title": "直生视频",
        "path": "工作流看板 - 直生视频",
        "image": "16_step_direct_video.png",
        "desc": "直生视频步骤基于首帧或首尾帧调用图生视频模型生成视频片段，并可拼接为完整视频。",
        "steps": ["用户进入直生视频步骤。", "系统展示视频模型、片段状态和生成入口。", "用户批量生成片段并在完成后拼接视频。"],
    },
    {
        "title": "项目资产库",
        "path": "项目 - 资产库",
        "image": "04_assets.png",
        "desc": "资产库统一管理项目生成的图片、视频、音频、文本和参考素材，支持预览和按类型查看。",
        "steps": ["用户进入项目资产库。", "系统读取项目关联资产。", "用户按资产类型查看、预览或下载项目成果。"],
    },
    {
        "title": "个人中心与充值",
        "path": "设置 - 充值",
        "image": "06_settings_recharge.png",
        "desc": "充值页面展示点数余额、充值记录和凭证上传入口，用于支持 AI 生成任务的点数消耗。",
        "steps": ["用户进入充值页面。", "系统展示当前点数和充值记录。", "用户提交充值金额和凭证，等待管理员审核。"],
    },
    {
        "title": "管理后台",
        "path": "管理员入口 - 用户统计",
        "image": "07_admin_users.png",
        "desc": "管理后台提供用户统计、充值审核、请求统计和项目查看能力，用于系统运营管理。",
        "steps": ["管理员进入后台页面。", "系统展示用户列表、点数和操作入口。", "管理员可查看用户项目、调整点数或审核充值订单。"],
    },
]


def set_font(run, font_name="宋体", size=12, bold=False, color=None):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph(p, line_spacing=1.5, first_line=False, after=6, before=0):
    fmt = p.paragraph_format
    fmt.line_spacing = line_spacing
    fmt.space_after = Pt(after)
    fmt.space_before = Pt(before)
    if first_line:
        fmt.first_line_indent = Cm(0.74)


def add_para(doc, text="", *, size=12, bold=False, align=None, first_line=True, after=6, before=0, font="宋体"):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    set_paragraph(p, first_line=first_line, after=after, before=before)
    run = p.add_run(text)
    set_font(run, font, size, bold)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    set_paragraph(p, first_line=False, after=10 if level == 1 else 6, before=12 if level == 1 else 8)
    size = 16 if level == 1 else 14 if level == 2 else 12
    run = p.add_run(text)
    set_font(run, "黑体", size, True)
    return p


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, *, bold=False):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if bold else WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph(p, line_spacing=1.3, first_line=False, after=0)
    p.clear()
    r = p.add_run(text)
    set_font(r, "宋体", 10.5, bold)


def add_table(doc, rows):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    header = table.rows[0].cells
    set_cell_text(header[0], "项目", bold=True)
    set_cell_text(header[1], "内容", bold=True)
    shade_cell(header[0], "EDEDED")
    shade_cell(header[1], "EDEDED")
    for key, value in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], key)
        set_cell_text(cells[1], value)
    doc.add_paragraph()
    return table


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"{SOFTWARE_NAME} 软件使用说明书 {VERSION}")
    set_font(r, "宋体", 9)


def add_cover(doc):
    for _ in range(7):
        doc.add_paragraph()
    add_para(doc, SOFTWARE_NAME, size=22, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, after=20, font="黑体")
    add_para(doc, "软件使用说明书", size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, after=44, font="黑体")
    for line in [
        f"版本号：{VERSION}",
        f"著作权人：{COPYRIGHT_OWNER}",
        f"完成日期：{COMPLETION_DATE}",
    ]:
        add_para(doc, line, size=16, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, after=12)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    set_paragraph(p, first_line=False, after=3)
    r = p.add_run(text)
    set_font(r, "宋体", 12)


def add_number(doc, number, text):
    p = doc.add_paragraph()
    set_paragraph(p, first_line=False, after=3)
    p.paragraph_format.left_indent = Cm(0.74)
    r = p.add_run(f"{number}. {text}")
    set_font(r, "宋体", 12)


def add_image(doc, image_path: Path, caption: str):
    if not image_path.exists():
        add_para(doc, f"图示：{caption}", align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(image_path), width=Cm(15.6))
    cap = add_para(doc, caption, size=9, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, after=8)
    cap.paragraph_format.keep_together = True


def add_toc_note(doc):
    add_para(doc, "目录", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, font="黑体")
    for item in [
        "第一章 软件概述",
        "第二章 软件安装与运行",
        "第三章 软件主要功能说明",
        "第四章 软件操作流程说明",
        "第五章 系统管理与维护",
        "第六章 常见问题与说明",
        "第七章 附录",
    ]:
        add_para(doc, item, first_line=False, after=3)


def build_docx():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(3.18)
    sec.right_margin = Cm(3.18)
    add_footer(sec)

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    add_cover(doc)
    doc.add_page_break()
    add_toc_note(doc)
    doc.add_page_break()

    add_heading(doc, "第一章 软件概述", 1)
    add_heading(doc, "1.1 软件简介", 2)
    add_para(doc, INTRO_TEXT)
    add_heading(doc, "1.2 运行环境", 2)
    add_para(doc, ENV_TEXT)
    add_heading(doc, "1.3 主要功能模块", 2)
    for text in [
        "创意扩散模块：根据用户输入的元构思生成多个创意方向，辅助用户确定项目切入角度。",
        "框架搭建模块：生成故事梗概、角色设定、幕结构、环境设定和视觉风格，并支持人工编辑。",
        "风格统一模块：生成并选择统一视觉风格样图，作为后续图像和视频生成的基准。",
        "人物设计模块：基于角色设定生成角色形象图，形成可复用的角色参考资产。",
        "概念图、宣传片、分镜、尾帧和直生视频模块：围绕影视前期制作流程生成视觉图像、视频片段和最终视频。",
        "资产库、充值与管理后台模块：提供项目资产沉淀、点数管理、充值审核和运营统计能力。",
    ]:
        add_bullet(doc, text)
    add_heading(doc, "1.4 软件特点", 2)
    add_para(doc, FEATURE_TEXT)

    add_heading(doc, "第二章 软件安装与运行", 1)
    add_heading(doc, "2.1 软件访问方式", 2)
    add_para(doc, "本软件采用 Web 应用形式运行，用户通过浏览器访问系统地址即可使用，无需在本地安装客户端程序。系统登录后按照项目维度管理创意内容、生成任务和产出资产。")
    add_heading(doc, "2.2 登录与权限", 2)
    add_para(doc, "系统通过账号登录方式识别用户身份。普通用户可管理本人项目、生成内容和资产；管理员可进入后台查看用户数据、审核充值记录并进行必要的运营管理。")

    add_heading(doc, "第三章 软件主要功能说明", 1)
    add_table(doc, [
        ("软件全称", SOFTWARE_NAME),
        ("软件简称", "AI Film Flow"),
        ("版本号", VERSION),
        ("软件类型", "Web 应用软件 / AI 影视创作工作流系统"),
        ("主要用途", "辅助用户从元构思开始，完成创意扩散、框架搭建、风格统一、人物设计、概念图、宣传片、分镜、尾帧和直生视频生成。"),
        ("运行方式", "用户通过浏览器访问系统，生成结果保存在项目资产库中。"),
    ])
    add_para(doc, "软件主要功能以项目为单位组织。用户创建项目后，系统按照工作流步骤依次推进，每一步均提供生成、预览、确认、编辑或重新生成等操作，使 AI 自动化生成和人工创意把控相结合。")

    add_heading(doc, "第四章 软件操作流程说明", 1)
    add_para(doc, "本章按照软件实际使用顺序列示主要页面、操作路径、操作步骤和运行界面。相关文字说明参考原软件使用说明书内容，并结合真实运行截图补充页面级说明。")
    for idx, item in enumerate(SECTIONS, 1):
        add_heading(doc, f"4.{idx} {item['title']}", 2)
        add_para(doc, f"操作路径：{item['path']}", bold=True, first_line=False)
        add_para(doc, f"功能说明：{item['desc']}")
        add_para(doc, "操作步骤：", bold=True, first_line=False, after=3)
        for step_no, step in enumerate(item["steps"], 1):
            add_number(doc, step_no, step)
        add_image(doc, SCREEN_DIR / item["image"], f"图 4-{idx} {item['title']}界面")

    add_heading(doc, "第五章 系统管理与维护", 1)
    add_para(doc, "系统管理功能主要面向管理员使用。管理员可查看用户列表、用户项目、点数情况、充值申请和系统运行数据，并根据实际运营需要对充值申请进行审核。")
    add_para(doc, "系统生成的图片、视频、音频、文本和参考素材统一保存在项目资产库中。用户可在项目内查看、预览或下载生成成果，管理员可通过后台辅助定位用户项目和使用情况。")

    add_heading(doc, "第六章 常见问题与说明", 1)
    for text in [
        "若生成任务长时间未完成，用户可刷新项目页面或在本地顺序生成工具中继续轮询任务状态。",
        "若某一图片或视频结果不符合预期，用户可在对应步骤中重新生成单项结果。",
        "若点数不足，用户需要进入充值页面提交充值申请并等待管理员审核通过。",
        "若同一项目包含多幕且镜头编号重复，系统会结合幕号和镜头编号共同定位镜头内容，保证分镜、首帧和视频片段对应一致。",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "第七章 附录", 1)
    add_para(doc, "本说明书所附界面图均对应软件主要功能页面，用于说明软件真实运行形态、主要操作流程和功能覆盖范围。页面内容与文字说明共同构成软件功能、操作方式和运行效果的说明材料。")

    doc.save(OUT_DOCX)
    return OUT_DOCX


def main():
    output = build_docx()
    print(output)


if __name__ == "__main__":
    main()
